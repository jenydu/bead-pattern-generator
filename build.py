"""Sync brand palettes from beadcolors and stage core/ + data/ into web/py for Pages."""
import csv
import json
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).parent
SRC = ROOT.parent / "beadcolors" / "raw"
DST = ROOT / "data" / "palettes"
WEB_PY = ROOT / "web" / "py"


# MARD "basic" chart: series A-H and M (221 codes). The code list matches the 221-color
# palette in github.com/Jett-Wu/Perler_Beads_Generator (MIT); RGB values stay beadcolors'.
MARD_221 = re.compile(r"[A-HM]\d+$")


DMC_XLSX = ROOT / "data" / "dmc_threadcolors.xlsx"  # made by tools/scrape_threadcolors.py


def build_dmc():
    """data/palettes/dmc.csv (ref, name, r, g, b, contributor) from the scraped workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("openpyxl not installed, keeping existing dmc.csv", file=sys.stderr)
        return
    if not DMC_XLSX.exists():
        return
    ws = load_workbook(DMC_XLSX, read_only=True, data_only=True)["DMC"]
    rows = [r for r in list(ws.iter_rows(values_only=True))[1:] if r[0] is not None]

    def key(row):  # B5200, Ecru, White first, then 1, 2, 3 ...
        n = str(row[0]).strip()
        return (1, int(n), "") if n.isdigit() else (0, 0, n)

    with open(DST / "dmc.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        for floss, name, r, g, b, *_ in sorted(rows, key=key):
            w.writerow([str(floss).strip(), name, r, g, b, "threadcolors.com"])


def sync_palettes():
    if not SRC.is_dir():
        print(f"{SRC} not found, keeping existing data/palettes", file=sys.stderr)
        return
    DST.mkdir(parents=True, exist_ok=True)
    for old in DST.glob("*.csv"):
        if old.stem != "dmc":  # generated from the workbook below, not copied from beadcolors
            old.unlink()
    for f in sorted(SRC.glob("*.csv")):
        if f.stem == "mard":
            continue  # shipped as mard_291 / mard_221
        shutil.copy2(f, DST / f.name)
    rows = (SRC / "mard.csv").read_text(encoding="utf-8").splitlines()
    (DST / "mard_291.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")
    basic = [r for r in rows if MARD_221.match(r.split(",")[0])]
    (DST / "mard_221.csv").write_text("\n".join(basic) + "\n", encoding="utf-8")
    build_dmc()
    lic = SRC.parent / "LICENSE"
    if lic.exists():
        shutil.copy2(lic, DST / "LICENSE-beadcolors")


def stage_web():
    """web/py mirrors the layout core/ expects: core/*.py and data/palettes/*.csv."""
    if WEB_PY.exists():
        shutil.rmtree(WEB_PY)
    shutil.copytree(ROOT / "core", WEB_PY / "core",
                    ignore=shutil.ignore_patterns("__pycache__"))
    shutil.copytree(DST, WEB_PY / "data" / "palettes")
    files = sorted(str(p.relative_to(WEB_PY)) for p in WEB_PY.rglob("*") if p.is_file())
    (WEB_PY / "manifest.json").write_text(json.dumps(files, indent=1))


if __name__ == "__main__":
    if "--no-sync" not in sys.argv:
        sync_palettes()
    stage_web()
    print("ok:", len(list(DST.glob("*.csv"))), "palettes")
