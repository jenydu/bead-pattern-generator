"""Scrape the DMC floss number -> RGB table from threadcolors.com into data/dmc_threadcolors.xlsx.

The whole table is in the front page's HTML, so this is a single request.
Usage: python tools/scrape_threadcolors.py   (needs openpyxl), then python build.py
"""
import html
import re
import ssl
import subprocess
import urllib.error
import urllib.request
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

URL = "https://threadcolors.com/"
OUT = Path(__file__).resolve().parent.parent / "data" / "dmc_threadcolors.xlsx"
ROW = re.compile(r'<tr><td style="background-color: #([0-9a-f]{6});"></td><td>([^<]*)</td>'
                 r'<td><a>([^<]*)</a></td><td>(\d+)</td><td>(\d+)</td><td>(\d+)</td>'
                 r'<td>([0-9a-f]{6})</td></tr>')


def fetch(url):
    req = urllib.request.Request(url, headers={"User-Agent": "bead-pattern-generator (one-off data import)"})
    try:
        return urllib.request.urlopen(req, timeout=30).read().decode("utf-8", "replace")
    except urllib.error.URLError as e:
        # Python may not trust a corporate proxy's certificate that the system does (curl uses
        # the system store). Certificate checking stays on either way.
        if not isinstance(e.reason, ssl.SSLError):
            raise
        return subprocess.run(["curl", "-sfL", "--max-time", "30", url], check=True,
                              capture_output=True).stdout.decode("utf-8", "replace")


def scrape():
    page = fetch(URL)
    rows = []
    for swatch, floss, name, r, g, b, hx in ROW.findall(page):
        r, g, b = int(r), int(g), int(b)
        if not (swatch == hx == "%02x%02x%02x" % (r, g, b)):
            raise ValueError(f"inconsistent row for floss {floss}")
        rows.append((html.unescape(floss).strip(), html.unescape(name).strip(), r, g, b, "#" + hx.upper()))
    if len({r[0] for r in rows}) != len(rows) or len(rows) < 400:
        raise ValueError(f"unexpected table: {len(rows)} rows")
    return rows


def write(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "DMC"
    ws.append(["Floss", "DMC name", "R", "G", "B", "Hex"])
    for i, row in enumerate(rows, start=2):
        ws.append(list(row))
        r, g, b = row[2:5]
        ws.cell(row=i, column=1).number_format = "@"
        ws.cell(row=i, column=6).fill = PatternFill("solid", fgColor=row[5][1:])
        ws.cell(row=i, column=6).font = Font(color="000000" if 0.299 * r + 0.587 * g + 0.114 * b > 140 else "FFFFFF")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, w in enumerate([9, 30, 6, 6, 6, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    notes = wb.create_sheet("Notes")
    for line in [f"DMC floss number to RGB, {len(rows)} threads.",
                 f"Source: {URL} (front page table), scraped {date.today().isoformat()}.",
                 "Re-create with: python tools/scrape_threadcolors.py",
                 "build.py reads sheet 'DMC' and writes data/palettes/dmc.csv.",
                 "The RGB values are threadcolors.com's screen approximation of each thread; "
                 "the real thread colors vary with the dye lot and the display."]:
        notes.append([line])
    notes.column_dimensions["A"].width = 100
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)


if __name__ == "__main__":
    rows = scrape()
    write(rows)
    print(f"wrote {len(rows)} threads to {OUT}")
